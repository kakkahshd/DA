"""gap_workbook — SATU berkas Excel berisi semua data yang masih harus diisi manusia,
sudah berisi baris datanya (model/SKU/material/akun) supaya user tidak mengisi dari nol.

Sheet yang bisa DIUNGGAH BALIK ke layar "Impor Harga · Rekening · BOM":
  MATERIAL · BOM_AKSESORIS · MODEL · REKENING · TOKO (importir lama)
  HARGA_JUAL_SKU · STOK_AWAL_FG · STOK_AWAL_MATERIAL · GAJI (ditambahkan di sini)
Sheet SALDO_AWAL & PIUTANG/HUTANG diunggah ke layar "Saldo Awal" (berkas yang sama boleh dipakai).
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone

import openpyxl
from openpyxl.styles import Font, PatternFill

from core.master_fill import MAT_COLS, REK_COLS, TOKO_COLS, BOM_COLS, MODEL_COLS, _head, _num
from core import opening_balance as _ob

SKU_PRICE_COLS = ["sku", "nama", "kode_model", "warna", "ukuran", "harga_saran_dari_saudara", "harga_jual", "keterangan"]
STOK_FG_COLS = ["sku", "nama", "kode_model", "warna", "ukuran", "lokasi_kode", "qty_awal", "hpp_satuan_sistem", "keterangan"]
STOK_MAT_COLS = ["kode", "nama", "tipe", "satuan_dasar", "lokasi_kode", "qty_awal", "harga_satuan_sistem", "keterangan"]
GAJI_COLS = ["kode_karyawan", "nama", "skema", "gaji_pokok_per_periode", "tarif_lembur_per_jam", "keterangan"]
DEFAULT_LOC_FG, DEFAULT_LOC_MAT = "GD-L1-RAK", "GD-L1"

PETUNJUK = [
    "DATA YANG MASIH HARUS DIISI — CV. Dewi Aditya (dibuat otomatis dari kondisi sistem saat ini)",
    "Baris yang tertulis di sini = data yang belum lengkap. Isi kolom KUNING saja; kolom lain hanya informasi.",
    "Baris yang kolom kuningnya dibiarkan kosong/0 TIDAK diubah — aman diunggah sebagian.",
    "",
    "UNGGAH ke Portal Keuangan → Master Akuntansi → Impor Harga · Rekening · BOM (berkas ini langsung):",
    "  MATERIAL          — material yang harganya masih 0: isi satuan_beli, isi_per_satuan_beli, harga_per_satuan_beli.",
    "  BOM_AKSESORIS     — aksesoris/bahan pendukung per model (benang, label, kancing…). Satu baris per bahan; kode_material lihat sheet REF_AKSESORIS.",
    "  MODEL             — berat_gram per model (untuk ongkir).",
    "  HARGA_JUAL_SKU    — SKU barang jadi yang belum berharga; harga_saran = harga SKU saudara di model yang sama (salin ke harga_jual bila setuju).",
    "  STOK_AWAL_FG      — stok fisik barang jadi per SKU per tanggal go-live (qty). Nilai rupiahnya diambil dari SALDO_AWAL (akun Persediaan), bukan dari sini.",
    "  STOK_AWAL_MATERIAL— stok fisik kain & aksesoris per tanggal go-live (qty dalam satuan dasar).",
    "  REKENING / TOKO   — no. rekening & atas nama; rekening pencairan tiap toko (sekarang default 1-1201 Bank BCA).",
    "  GAJI              — gaji pokok per periode & tarif lembur per karyawan (skema monthly/daily/piece).",
    "",
    "UNGGAH ke Portal Keuangan → Master Akuntansi → Saldo Awal (berkas yang sama):",
    "  SALDO_AWAL        — neraca penutup pembukuan lama per tanggal go-live; PIUTANG_*/HUTANG_* rinciannya per pelanggan/vendor.",
    "",
    "Tidak ada sandi di berkas ini. Techpack/foto/SOP diunggah per model di Portal RnD.",
]
YELLOW = PatternFill("solid", fgColor="FFF2CC")


def _mark(ws, cols: list[str], fill_cols: list[str]) -> None:
    idx = [cols.index(c) + 1 for c in fill_cols if c in cols]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in idx:
            row[i - 1].fill = YELLOW


def _widths(ws, widths):
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w


async def build_gap_workbook(db) -> tuple[bytes, dict]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PETUNJUK"
    for r in PETUNJUK:
        ws.append([r])
    ws.column_dimensions["A"].width = 130
    stats: dict = {}

    # ── MATERIAL: hanya yang harganya 0 (tanpa panel potongan) ──
    mats = await db.rahaza_materials.find({"active": {"$ne": False}, "type": {"$ne": "fg"}, "code": {"$not": {"$regex": "^CUT-"}}},
                                          {"_id": 0}).sort("code", 1).to_list(20000)
    ws = wb.create_sheet("MATERIAL")
    _head(ws, MAT_COLS)
    n = 0
    for m in mats:
        if float(m.get("unit_cost") or 0) > 0:
            continue
        ws.append([m.get("code"), m.get("name"), m.get("type"), m.get("category_name") or m.get("category"), m.get("unit"),
                   m.get("purchase_unit") or "", None, None, 0, m.get("min_stock") or None, "harga masih 0"])
        n += 1
    _mark(ws, MAT_COLS, ["satuan_beli", "isi_per_satuan_beli", "harga_per_satuan_beli"])
    _widths(ws, (14, 40, 10, 16, 12, 12, 16, 20, 20, 10, 24))
    stats["material_tanpa_harga"] = n

    # ── MODEL & BOM_AKSESORIS ──
    models = await db.rahaza_models.find({"active": {"$ne": False}}, {"_id": 0, "id": 1, "code": 1, "name": 1, "category_name": 1,
                                                                       "weight_gram": 1, "hpp": 1}).sort("code", 1).to_list(5000)
    has_bom, has_acc = set(), set()
    async for b in db.rahaza_boms.find({"active": {"$ne": False}, "is_active": True}, {"_id": 0, "model_id": 1, "materials": 1}):
        has_bom.add(b["model_id"])
        if any((ln.get("material_type") or "").lower() not in ("fabric", "") and not ln.get("is_cut_panel") for ln in b.get("materials") or []):
            has_acc.add(b["model_id"])
    ws = wb.create_sheet("MODEL")
    _head(ws, MODEL_COLS + ["hpp_sistem"])
    for m in models:
        ws.append([m["code"], m["name"], m.get("category_name"), float(m.get("weight_gram") or 0) or None,
                   "ya" if m["id"] in has_bom else "BELUM", "ya" if m["id"] in has_acc else "BELUM",
                   "" if m["id"] in has_bom else "belum punya BOM — isi BOM_AKSESORIS + kain di RnD", float(m.get("hpp") or 0) or None])
    _mark(ws, MODEL_COLS, ["berat_gram"])
    _widths(ws, (14, 30, 16, 12, 12, 16, 44, 14))
    stats["model_tanpa_berat"] = sum(1 for m in models if not m.get("weight_gram"))
    stats["model_tanpa_bom"] = sum(1 for m in models if m["id"] not in has_bom)
    stats["model_tanpa_aksesoris"] = sum(1 for m in models if m["id"] not in has_acc)

    ws = wb.create_sheet("BOM_AKSESORIS")
    _head(ws, BOM_COLS)
    for m in models:
        if m["id"] in has_acc:
            continue
        for _ in range(3):  # 3 baris kosong per model — tambah baris sendiri bila perlu
            ws.append([m["code"], m["name"], "", "", None, "", "isi kode_material (lihat REF_AKSESORIS) & qty_per_pcs"])
    _mark(ws, BOM_COLS, ["kode_material", "qty_per_pcs"])
    _widths(ws, (14, 30, 16, 36, 12, 10, 48))
    ws.freeze_panes = "C2"
    ws = wb.create_sheet("REF_AKSESORIS")
    _head(ws, ["kode_material", "nama", "tipe", "satuan_dasar", "harga_per_satuan_dasar"])
    for m in mats:
        if (m.get("type") or "").lower() != "fabric":
            ws.append([m.get("code"), m.get("name"), m.get("type"), m.get("unit"), float(m.get("unit_cost") or 0)])
    _widths(ws, (16, 44, 12, 12, 18))

    # ── HARGA_JUAL_SKU: FG tanpa harga + saran dari saudara ──
    fgs = await db.rahaza_materials.find({"type": "fg", "active": {"$ne": False}}, {"_id": 0}).sort("code", 1).to_list(20000)
    sib: dict = {}
    for f in fgs:
        p = float(f.get("retail_price_master") or 0)
        if p > 0:
            sib.setdefault(f.get("model_code"), set()).add(p)
    ws = wb.create_sheet("HARGA_JUAL_SKU")
    _head(ws, SKU_PRICE_COLS)
    n = 0
    for f in fgs:
        if float(f.get("retail_price_master") or 0) > 0:
            continue
        prices = sorted(sib.get(f.get("model_code")) or [])
        saran = prices[0] if len(prices) == 1 else None
        ket = "" if saran else ("model ini punya beberapa harga: " + " / ".join(f"{int(p):,}".replace(",", ".") for p in prices) if prices else "belum ada SKU berharga di model ini")
        ws.append([f.get("code"), f.get("name"), f.get("model_code"), f.get("color_name"), f.get("size_code"), saran, None,
                   ket or "salin harga_saran ke harga_jual bila setuju"])
        n += 1
    _mark(ws, SKU_PRICE_COLS, ["harga_jual"])
    _widths(ws, (26, 34, 12, 14, 10, 20, 14, 48))
    stats["sku_tanpa_harga"] = n

    # ── STOK_AWAL_FG & STOK_AWAL_MATERIAL ──
    ws = wb.create_sheet("STOK_AWAL_FG")
    _head(ws, STOK_FG_COLS)
    for f in fgs:
        ws.append([f.get("code"), f.get("name"), f.get("model_code"), f.get("color_name"), f.get("size_code"), DEFAULT_LOC_FG, None,
                   float(f.get("hpp") or 0) or None, ""])
    _mark(ws, STOK_FG_COLS, ["lokasi_kode", "qty_awal"])
    _widths(ws, (26, 34, 12, 14, 10, 12, 10, 16, 30))
    ws.freeze_panes = "B2"
    ws = wb.create_sheet("STOK_AWAL_MATERIAL")
    _head(ws, STOK_MAT_COLS)
    for m in mats:
        ws.append([m.get("code"), m.get("name"), m.get("type"), m.get("unit"), DEFAULT_LOC_MAT, None, float(m.get("unit_cost") or 0) or None, ""])
    _mark(ws, STOK_MAT_COLS, ["lokasi_kode", "qty_awal"])
    _widths(ws, (16, 44, 10, 12, 12, 10, 18, 30))
    ws.freeze_panes = "B2"
    ws = wb.create_sheet("REF_LOKASI")
    _head(ws, ["lokasi_kode", "nama", "tipe"])
    async for loc in db.rahaza_locations.find({"active": {"$ne": False}}, {"_id": 0, "code": 1, "name": 1, "type": 1}).sort("code", 1):
        ws.append([loc.get("code"), loc.get("name"), loc.get("type")])
    _widths(ws, (14, 30, 10))
    stats["sku_fg"], stats["material"] = len(fgs), len(mats)

    # ── REKENING & TOKO ──
    ws = wb.create_sheet("REKENING")
    _head(ws, REK_COLS)
    n = 0
    async for a in db.rahaza_cash_accounts.find({"active": {"$ne": False}}, {"_id": 0}).sort("coa_code", 1):
        ws.append([a.get("coa_code") or a.get("code"), a.get("name"), a.get("bank_name") or "", a.get("account_number") or "", a.get("account_holder") or ""])
        n += 1
    _mark(ws, REK_COLS, ["bank", "no_rekening", "atas_nama"])
    _widths(ws, (12, 34, 16, 22, 30))
    stats["rekening"] = n
    ws = wb.create_sheet("TOKO")
    _head(ws, TOKO_COLS + ["pic_email_sekarang"])
    async for t in db.marketing_platform_accounts.find({"status": {"$ne": "archived"}}, {"_id": 0}).sort("account_code", 1):
        ws.append([t.get("account_code"), t.get("account_name"), t.get("platform"), t.get("coa_cash_code") or "", t.get("pic_email") or t.get("pic_user_email") or ""])
    _mark(ws, TOKO_COLS, ["rekening_pencairan_kode_akun"])
    _widths(ws, (12, 30, 12, 28, 30))

    # ── GAJI ──
    ws = wb.create_sheet("GAJI")
    _head(ws, GAJI_COLS)
    emps = {e["id"]: e for e in await db.rahaza_employees.find({"active": {"$ne": False}}, {"_id": 0, "id": 1, "employee_code": 1, "name": 1}).to_list(5000)}
    n = 0
    async for p in db.rahaza_payroll_profiles.find({"active": {"$ne": False}}, {"_id": 0}):
        e = emps.get(p.get("employee_id")) or {}
        ws.append([e.get("employee_code"), e.get("name"), p.get("pay_scheme"), float(p.get("base_rate") or 0) or None,
                   float(p.get("overtime_rate") or 0) or None, "tarif lembur masih 0" if not p.get("overtime_rate") else ""])
        n += 1
    _mark(ws, GAJI_COLS, ["skema", "gaji_pokok_per_periode", "tarif_lembur_per_jam"])
    _widths(ws, (16, 30, 10, 22, 20, 30))
    stats["karyawan"] = n

    # ── SALDO_AWAL (+ rincian piutang/hutang) dari template resmi layar Saldo Awal ──
    ob_bytes, _ = await _ob.build_template(db)
    src = openpyxl.load_workbook(io.BytesIO(ob_bytes))
    for s in src.worksheets:
        if s.title == "PETUNJUK":
            continue
        ws = wb.create_sheet(s.title)
        for row in s.iter_rows(values_only=True):
            ws.append(list(row))
        for c in ws[1]:
            c.font = Font(bold=True)
        if s.title == "SALDO_AWAL":
            _mark(ws, ["kode_akun", "nama_akun", "tipe", "saldo_normal", "is_header", "debit", "kredit", "keterangan"], ["debit", "kredit"])
            _widths(ws, (12, 46, 12, 12, 10, 16, 16, 30))
        else:
            _widths(ws, (20, 30, 20, 16, 16, 30))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), stats


# ── PARSE & APPLY sheet tambahan (dipanggil dari master_fill) ────────────────────────
async def parse_extra_sheets(db, wb, errors: list) -> dict:
    out = {"sku_prices": [], "stock_fg": [], "stock_mat": [], "salaries": []}
    locs = {l["code"]: l for l in await db.rahaza_locations.find({}, {"_id": 0, "id": 1, "code": 1}).to_list(2000)}
    if "HARGA_JUAL_SKU" in wb.sheetnames:
        fg = {f["code"]: f for f in await db.rahaza_materials.find({"type": "fg"}, {"_id": 0, "id": 1, "code": 1, "retail_price_master": 1}).to_list(20000)}
        for i, r in enumerate(wb["HARGA_JUAL_SKU"].iter_rows(values_only=True, min_row=2), start=2):
            if not r or not r[0]:
                continue
            sku = str(r[0]).strip()
            try:
                harga = _num(r[6]) if len(r) > 6 else 0
            except ValueError:
                errors.append(f"HARGA_JUAL_SKU baris {i}: harga_jual bukan angka")
                continue
            if harga <= 0:
                continue
            if sku not in fg:
                errors.append(f"HARGA_JUAL_SKU baris {i}: SKU {sku} tidak ada")
                continue
            if abs(float(fg[sku].get("retail_price_master") or 0) - harga) > 1e-9:
                out["sku_prices"].append({"sku": sku, "id": fg[sku]["id"], "harga": harga})
    for sheet, key, coll_q, cols in (("STOK_AWAL_FG", "stock_fg", {"type": "fg"}, STOK_FG_COLS),
                                     ("STOK_AWAL_MATERIAL", "stock_mat", {"type": {"$ne": "fg"}}, STOK_MAT_COLS)):
        if sheet not in wb.sheetnames:
            continue
        items = {m["code"]: m for m in await db.rahaza_materials.find(coll_q, {"_id": 0, "id": 1, "code": 1, "unit": 1}).to_list(30000)}
        li, qi = cols.index("lokasi_kode"), cols.index("qty_awal")
        for i, r in enumerate(wb[sheet].iter_rows(values_only=True, min_row=2), start=2):
            if not r or not r[0]:
                continue
            code = str(r[0]).strip()
            try:
                qty = _num(r[qi]) if len(r) > qi else 0
            except ValueError:
                errors.append(f"{sheet} baris {i}: qty_awal bukan angka")
                continue
            if qty <= 0:
                continue
            loc = str(r[li]).strip() if len(r) > li and r[li] else ""
            if code not in items:
                errors.append(f"{sheet} baris {i}: kode {code} tidak ada di master")
                continue
            if loc not in locs:
                errors.append(f"{sheet} baris {i}: lokasi_kode '{loc}' tidak ada (lihat REF_LOKASI)")
                continue
            out[key].append({"code": code, "material_id": items[code]["id"], "location_id": locs[loc]["id"], "location_code": loc, "qty": qty, "unit": items[code].get("unit")})
    if "GAJI" in wb.sheetnames:
        emps = {e["employee_code"]: e for e in await db.rahaza_employees.find({}, {"_id": 0, "id": 1, "employee_code": 1}).to_list(5000)}
        profiles = {p["employee_id"]: p for p in await db.rahaza_payroll_profiles.find({}, {"_id": 0, "employee_id": 1, "base_rate": 1, "overtime_rate": 1, "pay_scheme": 1}).to_list(5000)}
        for i, r in enumerate(wb["GAJI"].iter_rows(values_only=True, min_row=2), start=2):
            if not r or not r[0]:
                continue
            code = str(r[0]).strip()
            if code not in emps:
                errors.append(f"GAJI baris {i}: kode_karyawan {code} tidak ada")
                continue
            try:
                base = _num(r[3]) if len(r) > 3 else 0
                ot = _num(r[4]) if len(r) > 4 else 0
            except ValueError:
                errors.append(f"GAJI baris {i}: gaji/tarif bukan angka")
                continue
            scheme = (str(r[2]).strip().lower() if len(r) > 2 and r[2] else "") or None
            if scheme and scheme not in ("monthly", "daily", "piece", "weekly"):
                errors.append(f"GAJI baris {i}: skema '{scheme}' tidak dikenal (monthly/daily/weekly/piece)")
                continue
            cur = profiles.get(emps[code]["id"]) or {}
            changed = (base > 0 and abs(base - float(cur.get("base_rate") or 0)) > 1e-9) or \
                      (ot > 0 and abs(ot - float(cur.get("overtime_rate") or 0)) > 1e-9) or \
                      (scheme and scheme != cur.get("pay_scheme"))
            if changed:
                out["salaries"].append({"employee_code": code, "employee_id": emps[code]["id"], "base": base, "ot": ot, "scheme": scheme})
    return out


async def apply_extra(db, parsed: dict, user: dict | None) -> dict:
    from core import stock_service
    now = datetime.now(timezone.utc)
    actor = {"id": str((user or {}).get("id") or "system"), "email": (user or {}).get("email", "")}
    for p in parsed.get("sku_prices") or []:
        await db.rahaza_materials.update_one({"id": p["id"]}, {"$set": {"retail_price_master": p["harga"], "updated_at": now}})
    n_stock = 0
    for key in ("stock_fg", "stock_mat"):
        for s in parsed.get(key) or []:
            already = await db.rahaza_material_movements.find_one({"type": "opening", "material_id": s["material_id"], "to_location_id": s["location_id"]}, {"_id": 0, "id": 1})
            if already:
                continue  # stok awal per item+lokasi hanya sekali — koreksi lewat Penyesuaian Stok
            await stock_service.add(s["material_id"], s["location_id"], s["qty"], ref={"source": "opening_stock", "reason": "Stok awal go-live (Excel)"}, actor=actor, db=db)
            await db.rahaza_material_movements.insert_one({"id": str(uuid.uuid4()), "created_at": now, "timestamp": now, "created_by": actor["id"],
                                                           "type": "opening", "material_id": s["material_id"], "qty": s["qty"], "from_location_id": None,
                                                           "to_location_id": s["location_id"], "ref_type": "opening_stock", "ref_id": None,
                                                           "notes": "Stok awal go-live dari Excel (nilai GL dari SALDO_AWAL)"})
            n_stock += 1
    for g in parsed.get("salaries") or []:
        upd = {"updated_at": now}
        if g["base"] > 0:
            upd["base_rate"] = g["base"]
        if g["ot"] > 0:
            upd["overtime_rate"] = g["ot"]
        if g["scheme"]:
            upd["pay_scheme"] = g["scheme"]
        await db.rahaza_payroll_profiles.update_one({"employee_id": g["employee_id"]}, {"$set": upd})
    return {"sku_prices_updated": len(parsed.get("sku_prices") or []), "opening_stock_rows": n_stock,
            "salaries_updated": len(parsed.get("salaries") or [])}
